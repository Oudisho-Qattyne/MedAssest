from rest_framework_simplejwt.views import TokenObtainPairView
from .serializers import CustomTokenObtainPairSerializer
from rest_framework.permissions import IsAdminUser
from rest_framework.decorators import api_view, permission_classes
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from users.models import User
from patients.models import Patient, FamilyMember
from aid_providers.models import AidProvider
from aid_requests.models import AidRequest, AidRequestType, AidRequestProvider

class CustomTokenObtainPairView(TokenObtainPairView):
    serializer_class = CustomTokenObtainPairSerializer

def _auto_width(ws, headers):
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(header)) + 4)


def _write_header(ws, headers):
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws.append(headers)
    for i in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _build_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title=title)
    _write_header(ws, headers)
    for row in rows:
        ws.append(row)
    _auto_width(ws, headers)
    return ws


@api_view(["GET"])
@permission_classes([IsAdminUser])
def export_excel(request):
    wb = Workbook()
    wb.remove(wb.active)

    # ── 1. Users ──
    users_headers = ["ID", "Email", "Full Name", "Role", "Is Active", "Created At"]
    users_rows = [
        [u.id, u.email, u.full_name, u.role, u.is_active, u.created_at.strftime("%Y-%m-%d %H:%M")]
        for u in User.objects.all()
    ]
    _build_sheet(wb, "Users", users_headers, users_rows)

    # ── 2. Patients ──
    patients_headers = [
        "ID", "First Name", "Middle Name", "Last Name", "Mother Name",
        "Birth Date", "Gender", "National Number", "Family Booklet No",
        "Residence", "Phone", "Telephone", "Marital Status", "Home Status",
        "Job Type", "Job Title", "Monthly Salary", "Special Needs",
        "Note", "Created By", "Created At",
    ]
    patients_rows = []
    for p in Patient.objects.all():
        patients_rows.append([
            p.id, p.first_name, p.middle_name, p.last_name, p.mother_full_name,
            p.birth_date.strftime("%Y-%m-%d") if p.birth_date else "",
            p.gender, p.national_number, p.family_booklet_no,
            p.current_residence, p.phone_number, p.telephone_number,
            p.marital_status, p.home_status, p.job_type, p.job_title,
            str(p.monthly_salary) if p.monthly_salary is not None else "",
            p.special_needs, p.note,
            p.created_by_user.email if p.created_by_user else "",
            p.created_at.strftime("%Y-%m-%d %H:%M"),
        ])
    _build_sheet(wb, "Patients", patients_headers, patients_rows)

    # ── 3. Family Members ──
    fm_headers = ["ID", "Patient ID", "Patient Name", "Full Name", "Relation", "Gender", "Birth Date", "Note"]
    fm_rows = []
    for fm in FamilyMember.objects.select_related("patient").all():
        fm_rows.append([
            fm.id, fm.patient_id, str(fm.patient), fm.full_name, fm.relation,
            fm.gender,
            fm.birth_date.strftime("%Y-%m-%d") if fm.birth_date else "",
            fm.note,
        ])
    _build_sheet(wb, "Family Members", fm_headers, fm_rows)

    # ── 4. Aid Request Types ──
    art_headers = ["ID", "Type Name", "Description"]
    art_rows = [[t.id, t.type_name, t.description] for t in AidRequestType.objects.all()]
    _build_sheet(wb, "Aid Request Types", art_headers, art_rows)

    # ── 5. Aid Providers ──
    prov_headers = ["ID", "Name", "Specialization", "Phone", "Email", "Category", "Created At"]
    prov_rows = []
    for p in AidProvider.objects.select_related("category").all():
        prov_rows.append([
            p.id, p.name, p.specialization, p.phone_number, p.email,
            p.category.category_name if p.category else "",
            p.created_at.strftime("%Y-%m-%d %H:%M"),
        ])
    _build_sheet(wb, "Aid Providers", prov_headers, prov_rows)

    # ── 6. Aid Requests ──
    req_headers = [
        "ID", "Patient ID", "Patient Name", "Status", "Request Type",
        "Description", "Estimated Cost", "Total Provided Amount",
        "Place of Aid", "Date of Aid",
    ]
    req_rows = []
    for r in AidRequest.objects.select_related("patient", "aid_request_type").prefetch_related("providers").all():
        total = 0
        estimated = r.estimated_cost or 0
        for link in r.providers.all():
            amt = link.aid_amount
            if amt is None:
                continue
            if link.type_of_aid_amount == "fixed":
                total += amt
            elif link.type_of_aid_amount == "percentage" and estimated > 0:
                total += (amt / 100) * estimated
        req_rows.append([
            r.id, r.patient_id, str(r.patient), r.request_status,
            r.aid_request_type.type_name if r.aid_request_type else "",
            r.description,
            float(r.estimated_cost) if r.estimated_cost is not None else "",
            round(total, 2),
            r.place_of_aid,
            r.date_of_aid.strftime("%Y-%m-%d") if r.date_of_aid else "",
        ])
    _build_sheet(wb, "Aid Requests", req_headers, req_rows)

    # ── 7. Aid Request Providers ──
    arp_headers = [
        "ID", "Request ID", "Patient Name", "Provider ID", "Provider Name",
        "Aid Type", "Aid Amount", "Amount Type", "Notes",
    ]
    arp_rows = []
    for arp in AidRequestProvider.objects.select_related("aid_request__patient", "aid_provider").all():
        arp_rows.append([
            arp.id, arp.aid_request_id, str(arp.aid_request.patient),
            arp.aid_provider_id, arp.aid_provider.name,
            arp.aid_type, float(arp.aid_amount) if arp.aid_amount is not None else "",
            arp.type_of_aid_amount, arp.notes,
        ])
    _build_sheet(wb, "Aid Request Providers", arp_headers, arp_rows)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="medassest_export.xlsx"'
    wb.save(response)
    return response
